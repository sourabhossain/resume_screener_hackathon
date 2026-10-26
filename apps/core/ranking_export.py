"""The job's master ranking as an Excel sheet, laid out like the recruiters' own."""
from copy import copy
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.models import Resume

COLUMNS = (
    ('Role', 22), ('Rank', 7), ('Candidate Name', 30), ('Overall Score (%)', 11), ('Verdict', 30),
    ('Years of Experience', 12), ('Current / Last Position', 36), ('Education (Highest Degree)', 32),
    ('Certifications', 30), ('Key Skills / Strengths', 60), ('Key Gaps / Missing Skills', 55),
    ('Recruiter Status', 18), ('Interview', 34),
)
LEFT_ALIGNED = {'Key Skills / Strengths', 'Key Gaps / Missing Skills', 'Interview'}
VERDICTS = {
    'interview': 'Strong Fit - Shortlist for Interview',
    'talent_pool': 'Moderate Fit - Backup / Secondary Review',
    'reject': 'Low Fit - Reject',
}
NAVY, WHITE = '1F3864', 'FFFFFF'
ROW_FILLS = {'interview': 'E2EFDA', 'talent_pool': 'FFF2CC', 'reject': 'FCE4E4'}
NEXT_ROUND, OUT = '00B050', 'FFC000'
LEGEND = (
    (ROW_FILLS['interview'], 'Strong Fit - shortlisted for interview'),
    (ROW_FILLS['talent_pool'], 'Moderate Fit - backup / secondary review'),
    (ROW_FILLS['reject'], 'Low Fit - reject'),
    (NEXT_ROUND, 'Name: interviewed and recommended for the next round'),
    (OUT, 'Name: withdrew, or not recommended after interview'),
)

THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type='solid')


def _text(items, sep='; '):
    return sep.join(str(i) for i in items if str(i).strip()) or '—'


def _interviews(resume):
    """One line per interview with a submitted evaluation, and the name highlight."""
    lines, highlight = [], None
    for interview in sorted(resume.interviews.all(), key=lambda i: (i.scheduled_date, i.phase)):
        evaluations = [e for e in interview.evaluations.all() if e.is_submitted]
        if not evaluations:
            continue
        parts = [interview.get_phase_display(), interview.scheduled_date.strftime('%d %b %Y')]
        average = interview.avg_score()
        if average is not None:
            parts.append(f'{average}/100')
        recommendations = list(dict.fromkeys(e.get_recommendation_display() for e in evaluations
                                             if e.recommendation))
        parts += recommendations
        lines.append(' · '.join(parts))
        verdicts = {e.recommendation for e in evaluations}
        if 'no' in verdicts:
            highlight = OUT
        elif 'yes' in verdicts:
            highlight = NEXT_ROUND
    if resume.recruiter_status == 'withdrawn':
        highlight = OUT
    return '\n'.join(lines) or '—', highlight


def ranked_rows(job):
    """Every scored candidate for the job, best first; equal scores share a rank."""
    resumes = list(Resume.objects.filter(job=job, screening_status='completed', final_score__isnull=False)
                   .only('candidate_name', 'final_score', 'recommendation', 'recruiter_status', 'screening_summary',
                         'experience_years', 'education', 'certifications', 'matched_skills', 'missing_skills')
                   .prefetch_related('interviews__evaluations').order_by('-final_score', 'candidate_name'))
    rows = []
    for index, resume in enumerate(resumes):
        # Sorted best first, so a candidate's rank is the place of the first
        # one with the same score.
        if not index or resume.final_score != resumes[index - 1].final_score:
            rank = index + 1
        summary = resume.screening_summary or {}
        years = resume.experience_years
        interview, highlight = _interviews(resume)
        rows.append({
            'resume': resume,
            'highlight': highlight,
            'values': [
                job.title,
                rank,
                resume.candidate_name,
                round(resume.final_score, 1),
                summary.get('verdict') or VERDICTS.get(resume.recommendation, '—'),
                summary.get('experience') or (f'{years:g} yrs' if years else '—'),
                summary.get('position') or '—',
                summary.get('education') or _text(resume.education),
                summary.get('certifications') or _text(resume.certifications),
                summary.get('strengths') or _text(resume.matched_skills, ', '),
                summary.get('gaps') or _text(resume.missing_skills, ', '),
                resume.get_recruiter_status_display() if resume.recruiter_status else 'New CV',
                interview,
            ],
        })
    return rows


def _put(cell, value):
    cell.value = value
    # openpyxl stores any string starting with "=" as a formula; text from a CV
    # or a sheet must stay text.
    if isinstance(value, str) and value.startswith('='):
        cell.data_type = 's'


def build_workbook(job) -> bytes:
    rows = ranked_rows(job)
    unscored = Resume.objects.filter(job=job).count() - len(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Master Ranking'
    last = get_column_letter(len(COLUMNS))

    ws.merge_cells(f'A1:{last}1')
    _put(ws['A1'], f'MASTER RANKING — {job.title.upper()}')
    ws['A1'].font = Font(name='Arial', size=14, bold=True, color=WHITE)
    ws['A1'].fill = _fill(NAVY)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    note = f'{len(rows)} candidates ranked by overall score · Downloaded {timezone.localtime():%d %b %Y, %I:%M %p}'
    if unscored:
        note += f' · {unscored} not yet scored, not listed'
    ws.merge_cells(f'A2:{last}2')
    _put(ws['A2'], note)
    ws['A2'].font = Font(name='Arial', size=9, italic=True, color='595959')
    ws['A2'].alignment = Alignment(horizontal='center')

    for col, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=3, column=col)
        _put(cell, title)
        cell.font = Font(name='Arial', size=10, bold=True, color=WHITE)
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[3].height = 32

    # openpyxl hashes every style object it is handed, which made styling the
    # cells most of a large sheet's time. Each distinct look is styled once and
    # its style record copied to the cells that share it.
    fonts = {bold: Font(name='Arial', size=10, bold=bold) for bold in (False, True)}
    aligns = {side: Alignment(horizontal=side, vertical='center', wrap_text=True) for side in ('left', 'center')}
    column_styles = [(fonts[title in ('Role', 'Candidate Name')], aligns['left' if title in LEFT_ALIGNED else 'center'])
                     for title, _ in COLUMNS]
    looks = {}

    def style(cell, col, fill_color):
        key = (col, fill_color)
        if key not in looks:
            cell.font, cell.alignment = column_styles[col - 1]
            cell.fill = _fill(fill_color)
            cell.border = BORDER
            if col == 4:
                cell.number_format = '0.0'
            looks[key] = copy(cell._style)
        else:
            cell._style = copy(looks[key])

    for index, row in enumerate(rows, start=4):
        row_fill = ROW_FILLS.get(row['resume'].recommendation, WHITE)
        for col, value in enumerate(row['values'], start=1):
            cell = ws.cell(row=index, column=col)
            _put(cell, value)
            style(cell, col, row['highlight'] if col == 3 and row['highlight'] else row_fill)

    end = 3 + len(rows)
    if rows:
        ws.conditional_formatting.add(
            f'D4:D{end}', ColorScaleRule(start_type='min', start_color='F8696B', mid_type='percentile',
                                         mid_value=50, mid_color='FFEB84', end_type='max', end_color='63BE7B'))
        ws.auto_filter.ref = f'A3:{last}{end}'

    _put(ws.cell(row=end + 2, column=1), 'Colour key')
    ws.cell(row=end + 2, column=1).font = Font(name='Arial', size=10, bold=True)
    for offset, (color, label) in enumerate(LEGEND, start=end + 3):
        ws.cell(row=offset, column=1).fill = _fill(color)
        ws.cell(row=offset, column=1).border = BORDER
        _put(ws.cell(row=offset, column=2), label)
        ws.cell(row=offset, column=2).font = Font(name='Arial', size=9)

    ws.freeze_panes = 'D4'
    ws.print_title_rows = '3:3'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
