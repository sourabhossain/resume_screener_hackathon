"""The job's master ranking downloaded as Excel."""
from datetime import date
from io import BytesIO

import openpyxl
import pytest
from django.urls import reverse

from apps.core.models import Resume


def _scored(job, name, score, recommendation, **extra):
    return Resume.objects.create(job=job, candidate_name=name, final_score=score, screening_status='completed',
                                 recommendation=recommendation, **extra)


def _sheet(client, job):
    response = client.get(reverse('core:job_ranking_xlsx', kwargs={'slug': job.slug}))
    assert response.status_code == 200
    return response, openpyxl.load_workbook(BytesIO(response.content)).active


def _rows(ws, count):
    return list(ws.iter_rows(min_row=4, max_row=3 + count, values_only=True))


@pytest.fixture
def ranked(sample_job, user):
    from apps.interviews.models import Interview, InterviewEvaluation
    asha = _scored(sample_job, 'Asha Rahman', 91.5, 'interview', recruiter_status='interviewing', screening_summary={
        'verdict': 'Strong Fit - Shortlist for Interview (Validate Sales Ownership)', 'experience': '16+ yrs',
        'position': 'Head of Sales, Example Ltd.', 'education': 'Not specified in CV', 'certifications': '—',
        'strengths': 'Merchant acquisition at scale', 'gaps': 'No payment-gateway exposure'})
    _scored(sample_job, 'Nabil Hasan', 70.6, 'talent_pool', education=['MBA-IBA'], certifications=['AML/CFT'],
            matched_skills=['B2B sales', 'Channel growth'], missing_skills=['Payments'], experience_years=12.5)
    _scored(sample_job, 'Rafiq Uddin', 54.6, 'reject', recruiter_status='withdrawn')
    _scored(sample_job, '=HYPERLINK("http://x")', 54.6, 'reject')
    Resume.objects.create(job=sample_job, candidate_name='Still Screening', screening_status='processing')
    interview = Interview.objects.create(resume=asha, scheduled_date=date(2026, 10, 12), status='completed')
    InterviewEvaluation.objects.create(interview=interview, evaluator=user, interviewer_name='Panel',
                                       recommendation='yes', is_submitted=True, scores={'adaptability': 4})
    return sample_job


@pytest.mark.django_db
def test_the_sheet_lists_scored_candidates_best_first_like_the_ranking(authenticated_client, ranked):
    response, ws = _sheet(authenticated_client, ranked)
    assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert response['Content-Disposition'].startswith('attachment; filename="') and '_ranking_' in response['Content-Disposition']
    assert ws['A1'].value == f'MASTER RANKING — {ranked.title.upper()}'
    assert '4 candidates ranked' in ws['A2'].value and '1 not yet scored' in ws['A2'].value
    assert [c.value for c in ws[3]][:5] == ['Role', 'Rank', 'Candidate Name', 'Overall Score (%)', 'Verdict']

    rows = _rows(ws, 4)
    assert [r[2] for r in rows][:2] == ['Asha Rahman', 'Nabil Hasan']
    assert [r[1] for r in rows] == [1, 2, 3, 3]
    assert 'Still Screening' not in [r[2] for r in ws.iter_rows(values_only=True)]


@pytest.mark.django_db
def test_the_sheet_wording_is_kept_and_ai_results_fill_the_rest(authenticated_client, ranked):
    _, ws = _sheet(authenticated_client, ranked)
    asha, nabil = _rows(ws, 2)
    assert asha[4:11] == ('Strong Fit - Shortlist for Interview (Validate Sales Ownership)', '16+ yrs',
                          'Head of Sales, Example Ltd.', 'Not specified in CV', '—',
                          'Merchant acquisition at scale', 'No payment-gateway exposure')
    assert asha[11] == 'Interviewing' and asha[12].startswith('Interview 1 · 12 Oct 2026')
    assert 'Yes - Next round' in asha[12]
    assert nabil[4:11] == ('Moderate Fit - Backup / Secondary Review', '12.5 yrs', '—', 'MBA-IBA', 'AML/CFT',
                           'B2B sales, Channel growth', 'Payments')


@pytest.mark.django_db
def test_rows_and_names_are_coloured_by_verdict_and_interview(authenticated_client, ranked):
    _, ws = _sheet(authenticated_client, ranked)
    fills = {ws.cell(row=r, column=3).value: (ws.cell(row=r, column=1).fill.start_color.rgb[-6:],
                                              ws.cell(row=r, column=3).fill.start_color.rgb[-6:]) for r in range(4, 8)}
    assert fills['Asha Rahman'] == ('E2EFDA', '00B050')
    assert fills['Nabil Hasan'] == ('FFF2CC', 'FFF2CC')
    assert fills['Rafiq Uddin'] == ('FCE4E4', 'FFC000')


@pytest.mark.django_db
def test_text_from_a_cv_is_never_stored_as_a_formula(authenticated_client, ranked):
    _, ws = _sheet(authenticated_client, ranked)
    cell = next(ws.cell(row=r, column=3) for r in range(4, 8)
                if str(ws.cell(row=r, column=3).value).startswith('=HYPERLINK'))
    assert cell.data_type == 's'


@pytest.mark.django_db
def test_signing_in_is_required(client, ranked):
    response = client.get(reverse('core:job_ranking_xlsx', kwargs={'slug': ranked.slug}))
    assert response.status_code == 302 and '/login' in response['Location']


@pytest.mark.django_db
def test_the_job_page_offers_the_download(authenticated_client, ranked):
    html = authenticated_client.get(reverse('core:job_detail', kwargs={'slug': ranked.slug})).content.decode()
    assert reverse('core:job_ranking_xlsx', kwargs={'slug': ranked.slug}) in html and 'Ranking (Excel)' in html
